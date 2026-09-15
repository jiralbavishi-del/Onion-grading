from django.contrib import admin
from django.http import HttpResponse
from django.urls import path
import openpyxl
from .models import Assessment, AssessmentImage, OnionVariety, DefectReferenceSample


@admin.register(OnionVariety)
class OnionVarietyAdmin(admin.ModelAdmin):
    list_display = ['name', 'code', 'variety_type', 'origin', 'order']
    search_fields = ['name', 'code', 'origin']


@admin.register(DefectReferenceSample)
class DefectReferenceSampleAdmin(admin.ModelAdmin):
    list_display = ['title', 'sample_id', 'tag', 'class_tag', 'order']
    search_fields = ['title', 'sample_id', 'tag']



class AssessmentImageInline(admin.TabularInline):
    model = AssessmentImage
    extra = 0
    readonly_fields = ['uploaded_at']


@admin.action(description="Export Selected as Excel")
def export_as_excel(modeladmin, request, queryset):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="grading_reports.xlsx"'

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = 'Assessments'

    columns = [
        'Lot ID', 'Created At', 'Supplier Name', 'Phone', 'State', 'Variety',
        'Grade', 'Size Class', 'Moisture %', 'Sprouting %', 'Damage %', 'Doubles %',
        'Inspector Name', 'Computed Score', 'Quality Status', 'Notes'
    ]
    
    for col_num, column_title in enumerate(columns, 1):
        worksheet.cell(row=1, column=col_num, value=column_title)

    for row_num, obj in enumerate(queryset, 2):
        row = [
            obj.lot_id, obj.created_at.replace(tzinfo=None) if obj.created_at else '',
            obj.supplier_name, obj.supplier_phone, obj.state, obj.variety,
            obj.grade, obj.size_class, obj.moisture, obj.sprouting, obj.damage, obj.doubles,
            obj.inspector_name, obj.computed_score, obj.quality_status, obj.notes
        ]
        for col_num, cell_value in enumerate(row, 1):
            worksheet.cell(row=row_num, column=col_num, value=cell_value)

    workbook.save(response)
    return response


@admin.register(Assessment)
class AssessmentAdmin(admin.ModelAdmin):
    actions = [export_as_excel]
    list_display = [
        'lot_id',
        'supplier_name',
        'variety',
        'grade',
        'computed_score',
        'quality_status',
        'created_at',
    ]
    list_filter = ['grade', 'quality_status', 'state', 'variety']
    search_fields = ['lot_id', 'supplier_name', 'supplier_phone', 'inspector_name']
    readonly_fields = ['computed_score', 'quality_status', 'created_at', 'updated_at']
    inlines = [AssessmentImageInline]

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('export-excel/', self.admin_site.admin_view(self.export_all_excel), name='grading_onionassessment_export_excel'),
        ]
        return custom_urls + urls

    def export_all_excel(self, request):
        # We can either export all or export based on current queryset. 
        # For simplicity and robust "download all" button, we export everything.
        queryset = Assessment.objects.all().order_by('-created_at')
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="all_grading_reports.xlsx"'

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Assessments'

        columns = [
            'Lot ID', 'Created At', 'Supplier Name', 'Phone', 'State', 'Variety',
            'Grade', 'Size Class', 'Moisture %', 'Sprouting %', 'Damage %', 'Doubles %',
            'Inspector Name', 'Computed Score', 'Quality Status', 'Notes'
        ]
        
        for col_num, column_title in enumerate(columns, 1):
            worksheet.cell(row=1, column=col_num, value=column_title)

        for row_num, obj in enumerate(queryset, 2):
            row = [
                obj.lot_id, obj.created_at.replace(tzinfo=None) if obj.created_at else '',
                obj.supplier_name, obj.supplier_phone, obj.state, obj.variety,
                obj.grade, obj.size_class, obj.moisture, obj.sprouting, obj.damage, obj.doubles,
                obj.inspector_name, obj.computed_score, obj.quality_status, obj.notes
            ]
            for col_num, cell_value in enumerate(row, 1):
                worksheet.cell(row=row_num, column=col_num, value=cell_value)

        workbook.save(response)
        return response
